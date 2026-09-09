#!/usr/bin/env python3
"""Deliverable 1 grading-data collector.

For every fork of the course template repo this script:

  1. Clones the fork with its FULL git history and all branches into
     ``../student_repos/<login>/`` (a sibling of this ``grading/`` folder),
     or fetches it if it is already there.
  2. Extracts the git history - every commit on every branch, with author,
     date, message, diff size, whether it is a merge, and a rough kind
     (feature / fix / docs / refactor / test / chore / other) - so you can
     see how many commits each group (and each member) made and what kind.
  3. Pulls the fork's issues from the GitHub API: open/closed state, who
     closed them, assignees, labels, comment authors and full comment text,
     task-list checkbox progress, and timeline event types - i.e. whether the
     students actually interacted with the assignment issues.
  4. Writes ``../student_repos/<login>/grading_data.json`` per fork and a
     combined ``summary.json`` + ``summary.csv`` in this folder.

The assignment issues are read from ``../issues_data.json`` so the summary
can report, per fork, which of the deliverable's issues exist / are open /
closed / have been touched.

This is a wrapper around the shared base code in
``../../populate.py`` (session handling, fork listing, pagination, rate
limiting, colour output). It only ever reads from GitHub - it never
creates, edits, or closes anything.

Usage (from this folder):

    python collect_grading_data.py --dry-run             # list forks, do nothing
    python collect_grading_data.py --test                # only the test student (krolabola)
    python collect_grading_data.py --fork-user alice --fork-user bob
    python collect_grading_data.py                       # every fork
    python collect_grading_data.py --skip-clone          # refresh issue data only

A GitHub token (``GITHUB_TOKEN`` / ``GH_TOKEN`` env var, or ``.env`` next to
``populate.py``, or ``--token``) is needed for the issue data and to avoid
low anonymous rate limits. Cloning public forks works without one.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import re
import subprocess
import sys
import time
from collections import Counter

# The shared base code (populate.py) lives two directories up, at the top
# level of projectmanager/, outside the per-deliverable folders.
BASE_DIR = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

import populate  # noqa: E402

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DELIVERABLE_DIR = os.path.dirname(SCRIPT_DIR)
DEFAULT_DATA_FILE = os.path.join(DELIVERABLE_DIR, "issues_data.json")
DEFAULT_REPOS_DIR = os.path.join(DELIVERABLE_DIR, "student_repos")
DEFAULT_TEST_FORK_USER = "krolabola"
DELIVERABLE_LABEL = "deliverable 1"
XLSX_NAME = "dprep-deliverable-1-grading.xlsx"

BOT_LOGINS = {"github-actions[bot]", "dependabot[bot]", "web-flow"}
UNIT = "\x1f"   # between fields of one commit
REC = "\x1e"    # before each commit record


# --------------------------------------------------------------------------- git

def _git(args: list[str], cwd: str | None = None, check: bool = True) -> str:
    env = {**os.environ, "GIT_TERMINAL_PROMPT": "0"}
    proc = subprocess.run(
        ["git", *args], cwd=cwd, env=env,
        capture_output=True, text=True, encoding="utf-8", errors="replace",
    )
    if check and proc.returncode != 0:
        raise RuntimeError(f"git {' '.join(args)} -> {proc.returncode}: {proc.stderr.strip()}")
    return proc.stdout


def clone_or_update(fork: dict, repos_dir: str) -> tuple[str, str]:
    """Clone the fork (full history, all branches) or fetch it if present.

    Returns (path, action) where action is 'cloned', 'updated', or 'reused'.
    """
    login = fork["owner"]["login"]
    dest = os.path.join(repos_dir, login)
    git_dir = os.path.join(dest, ".git")

    if os.path.isdir(git_dir):
        try:
            _git(["fetch", "--all", "--prune", "--tags", "--force"], cwd=dest)
            return dest, "updated"
        except RuntimeError as exc:
            print(populate.warn(f"    fetch failed, reusing existing clone: {exc}"))
            return dest, "reused"

    os.makedirs(repos_dir, exist_ok=True)
    # Default clone already retrieves every branch as a remote-tracking ref;
    # --no-single-branch is explicit about that. No --depth => full history.
    _git(["clone", "--no-single-branch", fork["clone_url"], dest])
    _git(["fetch", "--all", "--prune", "--tags", "--force"], cwd=dest)
    return dest, "cloned"


def has_readme(path: str) -> bool:
    """True if a top-level ``README.md`` exists in the clone (case-insensitive)."""
    try:
        return any(name.lower() == "readme.md" for name in os.listdir(path))
    except OSError:
        return False


def classify_commit(subject: str, is_merge: bool) -> str:
    if is_merge:
        return "merge"
    s = re.sub(r"^\W+", "", subject.strip().lower())
    words = set(s.split())
    table = [
        ("feature", ("feat", "feature", "add", "added", "implement", "implemented",
                     "create", "created", "new", "support")),
        ("fix", ("fix", "fixed", "fixes", "bug", "bugfix", "hotfix", "patch",
                 "resolve", "resolved", "correct", "corrected")),
        ("docs", ("doc", "docs", "readme", "documentation", "comment", "comments",
                  "typo", "wording")),
        ("refactor", ("refactor", "refactored", "cleanup", "rename", "renamed",
                      "restructure", "reorg", "tidy", "format", "lint")),
        ("test", ("test", "tests", "testing", "spec", "specs")),
        ("chore", ("chore", "wip", "update", "updated", "bump", "config", "setup",
                   "init", "initial", "misc", "gitignore", "merge")),
    ]
    for kind, keys in table:
        if s.startswith(keys) or words & set(keys):
            return kind
    return "other"


def collect_git(dest: str, default_branch: str) -> dict:
    """Walk every commit on every branch and summarise the history."""
    fmt = REC + UNIT.join(["%H", "%an", "%ae", "%aI", "%cI", "%P", "%D", "%s"])
    raw = _git(
        ["log", "--all", "--no-color", "--date-order", "--numstat",
         f"--pretty=format:{fmt}"],
        cwd=dest,
    )

    commits: list[dict] = []
    for record in raw.split(REC):
        record = record.strip("\n")
        if not record:
            continue
        lines = record.split("\n")
        fields = lines[0].split(UNIT)
        if len(fields) < 8:
            continue
        h, an, ae, authored, committed, parents, refs, subject = fields[:8]

        files = ins = dels = 0
        for ln in lines[1:]:
            parts = ln.split("\t")
            if len(parts) != 3:
                continue
            add, rem, _ = parts
            files += 1
            ins += 0 if add.strip() in ("", "-") else int(add)
            dels += 0 if rem.strip() in ("", "-") else int(rem)

        is_merge = len(parents.split()) > 1
        commits.append({
            "hash": h,
            "author_name": an,
            "author_email": ae.lower(),
            "authored_at": authored,
            "committed_at": committed,
            "is_merge": is_merge,
            "refs": [r.strip() for r in refs.split(",") if r.strip()],
            "subject": subject,
            "files_changed": files,
            "insertions": ins,
            "deletions": dels,
            "kind": classify_commit(subject, is_merge),
        })

    # Branches (local + remote-tracking). Skip symbolic refs - that is the
    # origin/HEAD pointer, which git prints as a bare remote name.
    branch_out = _git(
        ["for-each-ref", "--format=%(refname:short)\t%(symref)",
         "refs/heads", "refs/remotes"],
        cwd=dest, check=False,
    )
    branches = []
    for line in branch_out.splitlines():
        name, _, symref = line.partition("\t")
        name = name.strip()
        if not name or symref.strip() or name.endswith("/HEAD"):
            continue
        branches.append(name)
    branches = sorted(set(branches))
    leaf = lambda b: b.split("/")[-1]  # noqa: E731
    common = {default_branch, "main", "master", "HEAD"}
    extra_branches = sorted({leaf(b) for b in branches if leaf(b) not in common})

    authors: dict[str, dict] = {}
    for c in commits:
        a = authors.setdefault(
            c["author_email"],
            {"name": c["author_name"], "commits": 0, "insertions": 0, "deletions": 0},
        )
        a["commits"] += 1
        a["insertions"] += c["insertions"]
        a["deletions"] += c["deletions"]
    active_days = sorted({c["authored_at"][:10] for c in commits if c["authored_at"]})
    dates = [c["authored_at"] for c in commits if c["authored_at"]]

    return {
        "total_commits": len(commits),
        "merge_commits": sum(1 for c in commits if c["is_merge"]),
        "kinds": dict(Counter(c["kind"] for c in commits)),
        "authors": authors,
        "distinct_authors": len(authors),
        "active_days": len(active_days),
        "first_commit": min(dates) if dates else None,
        "last_commit": max(dates) if dates else None,
        "total_insertions": sum(c["insertions"] for c in commits),
        "total_deletions": sum(c["deletions"] for c in commits),
        "branches": branches,
        "extra_branches": extra_branches,
        "commits": commits,
    }


# ------------------------------------------------------------------------- issues

_CHECK_DONE = re.compile(r"^\s*[-*]\s+\[[xX]\]\s+", re.M)
_CHECK_TODO = re.compile(r"^\s*[-*]\s+\[ \]\s+", re.M)


def _checkboxes(body: str) -> dict:
    body = body or ""
    done = len(_CHECK_DONE.findall(body))
    todo = len(_CHECK_TODO.findall(body))
    return {"done": done, "total": done + todo}


def collect_issues(
    session, owner: str, repo: str, templates, base_owner: str
) -> tuple[dict, list[dict]]:
    """Return (issue summary, per-assignment-issue rows)."""
    api = populate.GITHUB_API
    issues = populate.get_open_and_closed_issues(session, owner, repo)
    non_student = {base_owner.lower()} | BOT_LOGINS

    items: list[dict] = []
    for iss in issues:
        num = iss["number"]

        comments = []
        if iss.get("comments", 0):
            comments = populate._get_paginated(
                session, f"{api}/repos/{owner}/{repo}/issues/{num}/comments"
            )
        comment_actors = sorted({
            c["user"]["login"] for c in comments if c.get("user")
        })
        comment_items = [
            {
                "author": (c.get("user") or {}).get("login"),
                "created_at": c.get("created_at"),
                "body": (c.get("body") or "").strip(),
            }
            for c in comments
        ]

        events = populate._get_paginated(
            session, f"{api}/repos/{owner}/{repo}/issues/{num}/events"
        )
        event_types = sorted({e["event"] for e in events})

        closed_by = (iss.get("closed_by") or {}).get("login")
        if not closed_by:
            for e in events:
                if e["event"] == "closed" and e.get("actor"):
                    closed_by = e["actor"]["login"]
                    break

        assignees = [a["login"] for a in iss.get("assignees", [])]
        boxes = _checkboxes(iss.get("body"))
        student_comment_actors = [a for a in comment_actors if a.lower() not in non_student]
        student_events = [
            e for e in events
            if (e.get("actor") or {}).get("login", "").lower() not in non_student
        ]
        interacted = bool(
            iss["state"] == "closed"
            or assignees
            or student_comment_actors
            or boxes["done"]
            or student_events
        )

        items.append({
            "number": num,
            "title": iss["title"],
            "state": iss["state"],
            "created_at": iss["created_at"],
            "closed_at": iss.get("closed_at"),
            "closed_by": closed_by,
            "opened_by": (iss.get("user") or {}).get("login"),
            "assignees": assignees,
            "labels": [l["name"] for l in iss.get("labels", [])],
            "comments_count": iss.get("comments", 0),
            "comment_actors": comment_actors,
            "comments": comment_items,
            "event_types": event_types,
            "checkboxes": boxes,
            "interacted": interacted,
        })

    def _is_student(login: str | None) -> bool:
        return bool(login) and login.lower() not in non_student

    summary = {
        "count": len(items),
        "open": sum(1 for it in items if it["state"] == "open"),
        "closed": sum(1 for it in items if it["state"] == "closed"),
        "closed_by_student": sum(1 for it in items if _is_student(it["closed_by"])),
        "with_comments": sum(1 for it in items if it["comments_count"]),
        "with_assignee": sum(1 for it in items if it["assignees"]),
        "with_labels": sum(
            1 for it in items
            if any(l.lower() != DELIVERABLE_LABEL for l in it["labels"])
        ),
        "interacted": sum(1 for it in items if it["interacted"]),
        "checkboxes_done": sum(it["checkboxes"]["done"] for it in items),
        "checkboxes_total": sum(it["checkboxes"]["total"] for it in items),
        "items": items,
    }

    by_title = {populate._normalize_title(it["title"]): it for it in items}
    assignment: list[dict] = []
    for t in templates:
        it = by_title.get(populate._normalize_title(t.title))
        assignment.append({
            "template_title": t.title,
            "found": it is not None,
            "number": it["number"] if it else None,
            "state": it["state"] if it else None,
            "interacted": bool(it["interacted"]) if it else False,
            "checkboxes": it["checkboxes"] if it else {"done": 0, "total": 0},
        })
    return summary, assignment


def collect_prs(session, owner: str, repo: str) -> dict:
    api = populate.GITHUB_API
    prs = populate._get_paginated(
        session, f"{api}/repos/{owner}/{repo}/pulls", params={"state": "all"}
    )
    items = []
    for pr in prs:
        base_repo = (pr["base"].get("repo") or {}).get("full_name")
        items.append({
            "number": pr["number"],
            "title": pr["title"],
            "state": pr["state"],
            "created_at": pr["created_at"],
            "merged_at": pr.get("merged_at"),
            "head": pr["head"]["ref"],
            "base": pr["base"]["ref"],
            "base_repo": base_repo,
            "to_upstream": bool(base_repo and base_repo.lower() != f"{owner}/{repo}".lower()),
            "user": (pr.get("user") or {}).get("login"),
        })
    return {
        "count": len(items),
        "open": sum(1 for p in items if p["state"] == "open"),
        "closed": sum(1 for p in items if p["state"] == "closed"),
        "merged": sum(1 for p in items if p["merged_at"]),
        "to_upstream": sum(1 for p in items if p["to_upstream"]),
        "items": items,
    }


# --------------------------------------------------------------------------- main

def _now() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")


def _write_csv(path: str, rows: list[dict]) -> None:
    with open(path, "w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=SUMMARY_COLUMNS)
        writer.writeheader()
        writer.writerows(rows)


def _write_json(path: str, payload: dict) -> None:
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(payload, fh, indent=2, ensure_ascii=False)


def _timestamped(path: str) -> str:
    root, ext = os.path.splitext(path)
    return f"{root}-{dt.datetime.now():%Y%m%d-%H%M%S}{ext}"


def _save_output(path: str, saver, *, label: str) -> tuple[str, bool]:
    """Run ``saver(path)``; if the file is locked (open in Excel), wait and
    retry, then fall back to a timestamped copy.

    Returns (path_written, is_fallback).
    """
    for attempt in range(3):
        try:
            saver(path)
            return path, False
        except (OSError, PermissionError) as exc:
            if attempt < 2:
                print(populate.warn(
                    f"  {label}: {os.path.basename(path)} is locked ({exc}); "
                    f"close it in Excel - retrying in 2s..."
                ))
                time.sleep(2)
            else:
                alt = _timestamped(path)
                print(populate.error(
                    f"  {label}: still locked - wrote {os.path.basename(alt)} instead."
                ))
                saver(alt)
                return alt, True


def summary_row(rec: dict) -> dict:
    """Flatten one fork record into a single wide row.

    Empty string ("") means "not collected" (that section was skipped or
    errored); a real 0 means "collected, and the count is zero".
    """
    fork = rec["fork"]
    git = rec.get("git")
    issues = rec.get("issues")
    prs = rec.get("pull_requests")
    assignment = rec.get("assignment_issues") or []

    row: dict = {
        "login": fork["owner"],
        "full_name": fork["full_name"],
        "clone_url": fork["clone_url"],
        "cloned": rec.get("git_action") or "no",
        "readme_present": (
            "" if rec.get("readme_present") is None
            else ("yes" if rec.get("readme_present") else "no")
        ),
        "git_error": rec.get("git_error") or "",
        "issues_error": rec.get("issues_error") or "",
    }

    # --- git history -----------------------------------------------------
    if git:
        kinds = git.get("kinds", {})
        authors = git.get("authors", {})
        row.update({
            "total_commits": git.get("total_commits", 0),
            "distinct_authors": git.get("distinct_authors", 0),
            "distinct_author_names": len({
                a["name"].strip().lower() for a in authors.values()
            }),
            "authors": "; ".join(
                f"{a['name']}({a['commits']})"
                for a in sorted(authors.values(), key=lambda x: -x["commits"])
            ),
            "active_days": git.get("active_days", 0),
            "first_commit": (git.get("first_commit") or "")[:10],
            "last_commit": (git.get("last_commit") or "")[:10],
            "merge_commits": git.get("merge_commits", 0),
            "feature_commits": kinds.get("feature", 0),
            "fix_commits": kinds.get("fix", 0),
            "docs_commits": kinds.get("docs", 0),
            "refactor_commits": kinds.get("refactor", 0),
            "test_commits": kinds.get("test", 0),
            "chore_commits": kinds.get("chore", 0),
            "other_commits": kinds.get("other", 0),
            "extra_branches": "; ".join(git.get("extra_branches", [])),
            "insertions": git.get("total_insertions", 0),
            "deletions": git.get("total_deletions", 0),
        })
    else:
        for key in (
            "total_commits", "distinct_authors", "distinct_author_names", "authors",
            "active_days", "first_commit", "last_commit", "merge_commits",
            "feature_commits", "fix_commits", "docs_commits", "refactor_commits",
            "test_commits", "chore_commits", "other_commits", "extra_branches",
            "insertions", "deletions",
        ):
            row[key] = ""

    # --- issue engagement (across ALL issues on the fork) ---------------
    if issues:
        row.update({
            "issues_total": issues.get("count", 0),
            "issues_open": issues.get("open", 0),
            "issues_closed": issues.get("closed", 0),
            "issues_closed_by_student": issues.get("closed_by_student", 0),
            "issues_with_comments": issues.get("with_comments", 0),
            "issues_with_assignee": issues.get("with_assignee", 0),
            "issues_with_labels": issues.get("with_labels", 0),
            "issues_interacted": issues.get("interacted", 0),
            "checkboxes_done": issues.get("checkboxes_done", 0),
            "checkboxes_total": issues.get("checkboxes_total", 0),
        })
    else:
        for key in (
            "issues_total", "issues_open", "issues_closed", "issues_closed_by_student",
            "issues_with_comments", "issues_with_assignee", "issues_with_labels",
            "issues_interacted", "checkboxes_done", "checkboxes_total",
        ):
            row[key] = ""

    # --- assignment issues (matched by title against issues_data.json) --
    if issues is not None:
        row.update({
            "assignment_found": sum(1 for a in assignment if a["found"]),
            "assignment_closed": sum(1 for a in assignment if a["state"] == "closed"),
            "assignment_interacted": sum(1 for a in assignment if a["interacted"]),
            "assignment_checkboxes_done": sum(a["checkboxes"]["done"] for a in assignment),
            "assignment_checkboxes_total": sum(a["checkboxes"]["total"] for a in assignment),
        })
    else:
        for key in (
            "assignment_found", "assignment_closed", "assignment_interacted",
            "assignment_checkboxes_done", "assignment_checkboxes_total",
        ):
            row[key] = ""

    # --- pull requests -------------------------------------------------
    if prs:
        row.update({
            "prs_total": prs.get("count", 0),
            "prs_open": prs.get("open", 0),
            "prs_merged": prs.get("merged", 0),
            "prs_to_upstream": prs.get("to_upstream", 0),
        })
    else:
        for key in ("prs_total", "prs_open", "prs_merged", "prs_to_upstream"):
            row[key] = ""

    return row


# Column order for summary.csv (also the key order in summary.json rows).
SUMMARY_COLUMNS = [
    "login", "full_name", "clone_url", "cloned", "readme_present",
    "total_commits", "distinct_authors", "distinct_author_names", "authors",
    "active_days", "first_commit", "last_commit", "merge_commits",
    "feature_commits", "fix_commits", "docs_commits", "refactor_commits",
    "test_commits", "chore_commits", "other_commits",
    "extra_branches", "insertions", "deletions",
    "issues_total", "issues_open", "issues_closed", "issues_closed_by_student",
    "issues_with_comments", "issues_with_assignee", "issues_with_labels",
    "issues_interacted", "checkboxes_done", "checkboxes_total",
    "assignment_found", "assignment_closed", "assignment_interacted",
    "assignment_checkboxes_done", "assignment_checkboxes_total",
    "prs_total", "prs_open", "prs_merged", "prs_to_upstream",
    "git_error", "issues_error",
]

# Metrics shown, in order, on the per-repo sheet's "Key metrics" block.
_REPO_METRICS = [c for c in SUMMARY_COLUMNS if c not in ("login", "full_name", "clone_url")]

_XL_BAD_CHARS = re.compile(r"[:\\/?*\[\]]")


def _sheet_name(login: str, used: set[str]) -> str:
    """A valid, unique (<=31 char) Excel sheet name for a fork."""
    name = _XL_BAD_CHARS.sub("_", login).strip() or "repo"
    name = name[:31]
    base, i = name, 2
    while name.lower() in {u.lower() for u in used}:
        suffix = f"_{i}"
        name = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(name)
    return name


def _cell_value(v):
    """Coerce a Python value into something Excel can store in one cell."""
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, (list, tuple)):
        return "; ".join(str(x) for x in v)
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    return v


def write_xlsx(path: str, records: list[dict], rows: list[dict],
               base_repo: str, collected_at: str) -> None:
    """Write the grading workbook: an Overview sheet, a Grading rubric sheet,
    plus one sheet per repository."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True)
    title_font = Font(bold=True, size=12, color="1F4E78")
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    wrap_top = Alignment(vertical="top", wrap_text=True)

    # Goodness fills for the grading cells on the Grading rubric sheet.
    rubric_fills = {
        "good": PatternFill("solid", fgColor="C6EFCE"),    # Yes / All / Strongly agree
        "okgood": PatternFill("solid", fgColor="E2EFDA"),  # Agree
        "some": PatternFill("solid", fgColor="FFEB9C"),    # Some / Neutral
        "warn": PatternFill("solid", fgColor="FCE4D6"),    # Disagree
        "bad": PatternFill("solid", fgColor="FFC7CE"),     # No / None / Strongly disagree
        "na": PatternFill("solid", fgColor="D9D9D9"),      # n/a
    }

    def autosize(ws, max_width=80):
        for col in ws.columns:
            letter = get_column_letter(col[0].column)
            longest = max((len(str(c.value)) for c in col if c.value is not None), default=0)
            ws.column_dimensions[letter].width = min(max(longest + 2, 10), max_width)

    def table(ws, start_row, title, headers, data_rows):
        r = start_row
        if title:
            c = ws.cell(row=r, column=1, value=title)
            c.font = title_font
            r += 1
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=r, column=ci, value=h)
            c.font = bold
            c.fill = head_fill
        r += 1
        for dr in data_rows:
            for ci, v in enumerate(dr, 1):
                ws.cell(row=r, column=ci, value=_cell_value(v)).alignment = wrap_top
            r += 1
        return r + 1  # leave a blank row after the block

    def _nsa(part, total, zero="None"):
        """None/Some/All (or zero/Some/All) for `part` out of `total`."""
        if not total:
            return "n/a"
        if not part:
            return zero
        return "All" if part >= total else "Some"

    def _rubric_auto_values(rec):
        """Auto-assessed rubric value per item for one fork, keyed by item."""
        issues = rec.get("issues")
        git = rec.get("git") or {}
        cloned = "Yes" if rec.get("git_action") in ("cloned", "updated", "reused") else "No"
        rp = rec.get("readme_present")
        readme = "n/a" if rp is None else ("Yes" if rp else "No")
        if issues:
            tot = issues.get("count", 0)
            closed_v = _nsa(issues.get("closed", 0), tot)
            cbx_v = "Yes" if issues.get("checkboxes_done", 0) else "No"
            labels_v = _nsa(issues.get("with_labels", 0), tot)
            assigned_v = _nsa(issues.get("with_assignee", 0), tot, zero="No")
            comments_v = _nsa(issues.get("with_comments", 0), tot)
        else:
            closed_v = cbx_v = labels_v = assigned_v = comments_v = "n/a"
        commits = git.get("total_commits")
        gt5 = "n/a" if commits is None else ("Yes" if commits > 5 else "No")
        gt10 = "n/a" if commits is None else ("Yes" if commits > 10 else "No")
        return {
            "Repository is cloned": cloned,
            "Issues are closed": closed_v,
            "Checkboxes are used": cbx_v,
            "README.md is present": readme,
            "Labels are added to issues": labels_v,
            "Issues are assigned": assigned_v,
            "Comments are given to issues": comments_v,
            "Amount of commits higher than 5": gt5,
            "Amount of commits higher than 10": gt10,
        }

    # (item, scale options) in display order.
    _RUBRIC_AUTO_ITEMS = [
        ("Repository is cloned",         ["No", "Yes"]),
        ("Issues are closed",            ["None", "Some", "All"]),
        ("Checkboxes are used",          ["No", "Yes"]),
        ("README.md is present",         ["No", "Yes"]),
        ("Labels are added to issues",   ["None", "Some", "All"]),
        ("Issues are assigned",          ["No", "Some", "All"]),
        ("Comments are given to issues", ["None", "Some", "All"]),
    ]
    _LIKERT = ["Strongly disagree", "Disagree", "Neutral", "Agree", "Strongly agree"]
    _RUBRIC_MANUAL_ITEMS = [
        ("Commit history is evenly distributed over time", _LIKERT),
        ("Commit history is evenly distributed across people", _LIKERT),
    ]
    # Section 2 "Coding" - all manually graded, no auto-collected data.
    _RUBRIC_CODING_ITEMS = [
        ("qmd file is created", ["No", "Yes"]),
        ("qmd file runs following the instructions", ["No", "Yes"]),
        ("Clean summary is created", _LIKERT),
        ("Summary output contains deep insights", _LIKERT),
    ]
    # Section 3 "Versioning" - an auto block (from the git history) + a manual item.
    _RUBRIC_VERSIONING_AUTO_ITEMS = [
        ("Amount of commits higher than 5", ["No", "Yes"]),
        ("Amount of commits higher than 10", ["No", "Yes"]),
    ]
    _RUBRIC_VERSIONING_MANUAL_ITEMS = [
        ("Quality of commits is high", _LIKERT),
    ]
    # Score weight per grading value; anything else (incl. blank) scores 0.
    _RUBRIC_SCORE_WEIGHTS = [
        (1.0, ("Yes", "All", "Strongly agree")),
        (0.75, ("Agree",)),
        (0.5, ("Some", "Neutral")),
        (0.25, ("Disagree",)),
    ]
    _RUBRIC_SCORE_NOTE = (
        "Yes/All/Strongly agree=1, Agree=0.75, Some/Neutral=0.5, Disagree=0.25, "
        "No/None/Strongly disagree=0; n/a & blank excluded"
    )
    # Fill for each grading value, worst -> best, used for conditional formatting.
    _RUBRIC_VALUE_FILLS = [
        ("Strongly disagree", rubric_fills["bad"]),
        ("Disagree", rubric_fills["warn"]),
        ("Neutral", rubric_fills["some"]),
        ("Agree", rubric_fills["okgood"]),
        ("Strongly agree", rubric_fills["good"]),
        ("No", rubric_fills["bad"]),
        ("None", rubric_fills["bad"]),
        ("Some", rubric_fills["some"]),
        ("Yes", rubric_fills["good"]),
        ("All", rubric_fills["good"]),
        ("n/a", rubric_fills["na"]),
    ]

    def rubric_matrix(ws, start_row, records):
        """Deliverable 1 grading rubric as an items x repos matrix.

        Written on its own sheet (2nd in the workbook): one row per rubric
        item, one column per fork, grouped into numbered categories
        ("1. Project management", "2. Coding", "3. Versioning"), each split
        into an "Automatic grading" and/or "Manual grading" sub-block.
        Automatic items are pre-filled per fork from the collected data; every
        assessment cell carries a dropdown so a grader can override it. Manual
        items are left blank with a dropdown.

        Every assessment cell is coloured by "goodness" with conditional
        formatting (green best -> red worst, grey n/a) so it follows a grader's
        override. Every sub-block gets its own live "Score" row per fork -
        Yes/All/Strongly agree=1, Agree=0.75, Some/Neutral=0.5, Disagree=0.25,
        others=0, n/a & blank excluded - shown as a percentage and
        colour-scaled red -> green.
        """
        from openpyxl.worksheet.datavalidation import DataValidation

        logins = [rec["fork"]["owner"] for rec in records]
        auto_values = [_rubric_auto_values(rec) for rec in records]
        blank_values = [None] * len(records)
        n = len(logins)
        first_col = get_column_letter(3)
        last_col = get_column_letter(2 + max(n, 1))

        r = start_row

        def _paint(first_row, last_row):
            """Colour every value cell in rows first_row..last_row by goodness."""
            if not n:
                return
            rng = f"{first_col}{first_row}:{last_col}{last_row}"
            for value, fill in _RUBRIC_VALUE_FILLS:
                ws.conditional_formatting.add(rng, CellIsRule(
                    operator="equal", formula=[f'"{value}"'], fill=fill,
                ))

        def _heading(text):
            nonlocal r
            ws.cell(row=r, column=1, value=text).font = bold
            r += 1
            for ci, h in enumerate(["item", "scale", *logins], 1):
                cell = ws.cell(row=r, column=ci, value=h)
                cell.font = bold
                cell.fill = head_fill
            r += 1

        def _score_row(first, last):
            """A per-fork "Score" row (weighted %, colour-scaled) for a block."""
            nonlocal r
            ws.cell(row=r, column=1, value="Score").font = bold
            ws.cell(row=r, column=2, value=_RUBRIC_SCORE_NOTE).alignment = wrap_top
            for j in range(n):
                col = get_column_letter(3 + j)
                cr = f"{col}{first}:{col}{last}"
                denom = f'SUMPRODUCT(--({cr}<>"n/a"),--({cr}<>""))'
                terms = []
                for w, values in _RUBRIC_SCORE_WEIGHTS:
                    counts = "+".join(f'SUMPRODUCT(--({cr}="{v}"))' for v in values)
                    if w == 1.0:
                        terms.append(counts)
                    elif len(values) == 1:
                        terms.append(f"{w}*{counts}")
                    else:
                        terms.append(f"{w}*({counts})")
                pts = "+".join(terms)
                c = ws.cell(row=r, column=3 + j,
                            value=f'=IF({denom}=0,"n/a",({pts})/{denom})')
                c.number_format = "0%"
                c.font = bold
            ws.conditional_formatting.add(
                f"{first_col}{r}:{last_col}{r}",
                ColorScaleRule(
                    start_type="num", start_value=0, start_color="F8696B",
                    mid_type="num", mid_value=0.5, mid_color="FFEB84",
                    end_type="num", end_value=1, end_color="63BE7B",
                ),
            )
            r += 2

        def _section(label, items, values_by_repo):
            nonlocal r
            ws.cell(row=r, column=1, value=label).font = title_font
            r += 1
            first = r
            for item, opts in items:
                ws.cell(row=r, column=1, value=item).alignment = wrap_top
                ws.cell(row=r, column=2, value=" / ".join(opts)).alignment = wrap_top
                dv = DataValidation(
                    type="list", formula1='"%s"' % ",".join(opts), allow_blank=True
                )
                ws.add_data_validation(dv)
                for j, values in enumerate(values_by_repo):
                    cell = ws.cell(row=r, column=3 + j,
                                   value=(values or {}).get(item, ""))
                    cell.alignment = wrap_top
                    dv.add(cell)
                r += 1
            last = r - 1
            _paint(first, last)
            r += 1  # blank row between the items and the score
            if n:
                _score_row(first, last)
            return first, last

        _heading("1. Project management")
        _section("Automatic grading", _RUBRIC_AUTO_ITEMS, auto_values)
        _section("Manual grading", _RUBRIC_MANUAL_ITEMS, blank_values)

        r += 1
        _heading("2. Coding")
        _section("Manual grading", _RUBRIC_CODING_ITEMS, blank_values)

        r += 1
        _heading("3. Versioning")
        _section("Automatic grading", _RUBRIC_VERSIONING_AUTO_ITEMS, auto_values)
        _section("Manual grading", _RUBRIC_VERSIONING_MANUAL_ITEMS, blank_values)
        return r

    def repo_charts(ws, start_row, git):
        """Per-repo sheet charts: a day-distribution bar (commits per day,
        stacked by author) and a user-distribution pie (commits per author).

        Source data is written into columns A-B below the sheet's tables; the
        charts float to the right.
        """
        from openpyxl.chart import BarChart, PieChart, Reference

        commits = git.get("commits") or []
        authors_map = git.get("authors") or {}
        if not commits or not authors_map:
            return start_row

        ordered = sorted(authors_map.items(), key=lambda kv: -kv[1]["commits"])
        emails = [e for e, _ in ordered]
        names = [a["name"] for _, a in ordered]
        col_of = {e: i for i, e in enumerate(emails)}

        per_day: dict[str, list[int]] = {}
        for c in commits:
            day = (c.get("authored_at") or "")[:10]
            i = col_of.get(c["author_email"])
            if not day or i is None:
                continue
            per_day.setdefault(day, [0] * len(names))[i] += 1
        days = sorted(per_day)

        anchor_col = get_column_letter(max(16, len(names) + 3))
        r = start_row
        ws.cell(row=r, column=1, value="Charts").font = title_font
        r += 1

        # --- stacked bar: commits per day, one series per author ---
        hdr = r
        ws.cell(row=hdr, column=1, value="commits/day").font = bold
        for j, nm in enumerate(names):
            cell = ws.cell(row=hdr, column=2 + j, value=nm)
            cell.font = bold
            cell.fill = head_fill
        r += 1
        first = r
        for day in days:
            ws.cell(row=r, column=1, value=day)
            for j, v in enumerate(per_day[day]):
                ws.cell(row=r, column=2 + j, value=v)
            r += 1
        last = r - 1

        if days:
            bar = BarChart()
            bar.type = "col"
            bar.grouping = "stacked"
            bar.overlap = 100
            bar.title = "Commit history - commits per day, stacked by author"
            bar.y_axis.title = "commits"
            bar.add_data(
                Reference(ws, min_col=2, max_col=1 + len(names), min_row=hdr, max_row=last),
                titles_from_data=True,
            )
            bar.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
            bar.height, bar.width = 8, 24
            ws.add_chart(bar, f"{anchor_col}{start_row}")

        r += 1

        def _author_pie(label_row_value, value_fn, title, anchor_row):
            nonlocal r
            top = r
            ws.cell(row=top, column=1, value="author").font = bold
            ws.cell(row=top, column=2, value=label_row_value).font = bold
            r += 1
            body = r
            for _, a in ordered:
                ws.cell(row=r, column=1, value=a["name"])
                ws.cell(row=r, column=2, value=value_fn(a))
                r += 1
            pie = PieChart()
            pie.title = title
            pie.add_data(Reference(ws, min_col=2, min_row=top, max_row=r - 1),
                         titles_from_data=True)
            pie.set_categories(Reference(ws, min_col=1, min_row=body, max_row=r - 1))
            pie.height, pie.width = 8, 12
            ws.add_chart(pie, f"{anchor_col}{anchor_row}")
            r += 1

        _author_pie("commits", lambda a: a["commits"],
                    "Commits per author", start_row + 18)
        return r

    wb = Workbook()

    # ---------------------------------------------------------------- Overview
    ov = wb.active
    ov.title = "Overview"
    ov.cell(row=1, column=1, value=f"Grading overview - {base_repo}").font = title_font
    ov.cell(row=2, column=1, value=f"collected {collected_at}  |  {len(rows)} fork(s)")
    header_row = 4
    for ci, col in enumerate(SUMMARY_COLUMNS, 1):
        c = ov.cell(row=header_row, column=ci, value=col)
        c.font = bold
        c.fill = head_fill
    for ri, row in enumerate(rows, header_row + 1):
        for ci, col in enumerate(SUMMARY_COLUMNS, 1):
            ov.cell(row=ri, column=ci, value=_cell_value(row.get(col, "")))
    last_col = get_column_letter(len(SUMMARY_COLUMNS))
    ov.auto_filter.ref = f"A{header_row}:{last_col}{header_row + len(rows)}"
    ov.freeze_panes = f"C{header_row + 1}"
    autosize(ov, max_width=45)

    # ----------------------------------------------------- Grading rubric sheet
    rb = wb.create_sheet("Grading rubric", 1)  # 2nd sheet, after Overview
    rb.cell(row=1, column=1, value=f"Grading rubric - {base_repo}").font = title_font
    rb.cell(row=2, column=1, value=f"collected {collected_at}  |  {len(rows)} fork(s)")
    rubric_matrix(rb, 4, records)
    autosize(rb, max_width=45)

    # --------------------------------------------------------- per-repo sheets
    used = {"Overview", "Grading rubric"}
    for rec, row in zip(records, rows):
        fork = rec["fork"]
        ws = wb.create_sheet(_sheet_name(fork["owner"], used))
        git = rec.get("git") or {}
        issues = rec.get("issues") or {}
        prs = rec.get("pull_requests") or {}
        assignment = rec.get("assignment_issues") or []

        ws.cell(row=1, column=1, value=fork["full_name"]).font = title_font
        r = table(ws, 3, "Repository", ["field", "value"], [
            ("html_url", fork.get("html_url")),
        ])

        if git.get("authors"):
            r = table(ws, r, "Authors", ["name", "email", "commits", "lines +", "lines -"], [
                (a["name"], email, a["commits"], a.get("insertions", 0), a.get("deletions", 0))
                for email, a in sorted(git["authors"].items(), key=lambda kv: -kv[1]["commits"])
            ])

        if git.get("branches"):
            extra = set(git.get("extra_branches", []))
            r = table(ws, r, "Branches", ["branch", "beyond main/default"],
                      [(b, "yes" if b.split("/")[-1] in extra else "") for b in git["branches"]])

        if git.get("commits"):
            r = table(
                ws, r, f"Commits ({len(git['commits'])})",
                ["authored", "author", "email", "kind", "merge",
                 "files", "insertions", "deletions", "subject", "refs"],
                [
                    (c["authored_at"][:16].replace("T", " "), c["author_name"], c["author_email"],
                     c["kind"], c["is_merge"], c["files_changed"], c["insertions"],
                     c["deletions"], c["subject"], c["refs"])
                    for c in git["commits"]
                ],
            )

        if issues.get("items"):
            r = table(
                ws, r, f"Issues ({issues.get('count', 0)})",
                ["#", "title", "state", "opened_by", "closed_by", "closed_at",
                 "assignees", "labels", "comments", "comment_actors",
                 "cb_done", "cb_total", "events", "interacted"],
                [
                    (it["number"], it["title"], it["state"], it["opened_by"], it["closed_by"],
                     it["closed_at"], it["assignees"], it["labels"], it["comments_count"],
                     it["comment_actors"], it["checkboxes"]["done"], it["checkboxes"]["total"],
                     it["event_types"], it["interacted"])
                    for it in issues["items"]
                ],
            )

            comment_rows = [
                (it["number"], cm.get("author"), cm.get("created_at"), cm.get("body"))
                for it in issues["items"]
                for cm in it.get("comments", [])
            ]
            if comment_rows:
                r = table(
                    ws, r, f"Issue comments ({len(comment_rows)})",
                    ["#", "author", "created_at", "body"], comment_rows,
                )

        if assignment:
            r = table(
                ws, r, "Assignment issues (matched from issues_data.json)",
                ["template_title", "found", "#", "state", "interacted", "cb_done", "cb_total"],
                [
                    (a["template_title"], a["found"], a["number"], a["state"],
                     a["interacted"], a["checkboxes"]["done"], a["checkboxes"]["total"])
                    for a in assignment
                ],
            )

        if prs.get("items"):
            r = table(
                ws, r, f"Pull requests ({prs.get('count', 0)})",
                ["#", "title", "state", "head", "base", "base_repo",
                 "to_upstream", "created_at", "merged_at", "user"],
                [
                    (p["number"], p["title"], p["state"], p["head"], p["base"], p["base_repo"],
                     p["to_upstream"], p["created_at"], p["merged_at"], p["user"])
                    for p in prs["items"]
                ],
            )

        # Key metrics live at the bottom of the sheet.
        r = table(ws, r, "Key metrics", ["metric", "value"],
                  [(m, row.get(m, "")) for m in _REPO_METRICS])

        repo_charts(ws, r, git)

        autosize(ws, max_width=70)

    wb.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(
        description=f"Collect {DELIVERABLE_LABEL} grading data (git history + issues) "
                    f"from every fork."
    )
    parser.add_argument("--owner", default=populate.DEFAULT_OWNER,
                        help="Owner of the base repository whose forks are graded.")
    parser.add_argument("--repo", default=populate.DEFAULT_REPO,
                        help="Name of the base repository.")
    parser.add_argument("--data-file", default=DEFAULT_DATA_FILE,
                        help="Assignment issue templates (default: ../issues_data.json).")
    parser.add_argument("--repos-dir", default=DEFAULT_REPOS_DIR,
                        help="Where forks are cloned (default: ../student_repos).")
    parser.add_argument("--out-dir", default=SCRIPT_DIR,
                        help=f"Where summary.json / summary.csv / {XLSX_NAME} are written "
                             "(default: this folder).")
    parser.add_argument("--no-xlsx", action="store_true",
                        help=f"Skip the {XLSX_NAME} workbook (needs openpyxl).")
    parser.add_argument("--fork-user", action="append", default=None,
                        help="Restrict to this GitHub login (repeatable).")
    parser.add_argument("--test", action="store_true",
                        help=f"Shortcut for --fork-user {DEFAULT_TEST_FORK_USER} "
                             f"(the test student), unless --fork-user is given.")
    parser.add_argument("--token",
                        default=os.environ.get("GITHUB_TOKEN") or os.environ.get("GH_TOKEN"),
                        help="GitHub token (defaults to GITHUB_TOKEN/GH_TOKEN env var).")
    parser.add_argument("--limit", type=int, default=None,
                        help="Only process the first N forks (debugging).")
    parser.add_argument("--skip-clone", action="store_true",
                        help="Do not clone/fetch; use whatever is already in --repos-dir.")
    parser.add_argument("--skip-git", action="store_true",
                        help="Do not read git history.")
    parser.add_argument("--skip-issues", action="store_true",
                        help="Do not query the issues API.")
    parser.add_argument("--dry-run", action="store_true",
                        help="List the forks that would be processed, then stop.")
    args = parser.parse_args()

    repos_dir = os.path.abspath(args.repos_dir)
    out_dir = os.path.abspath(args.out_dir)

    if not args.token and not args.skip_issues:
        print(populate.warn(
            "Warning: no GitHub token (set GITHUB_TOKEN or --token). Issue data will "
            "likely fail and rate limits are low. Cloning still works."
        ), file=sys.stderr)

    templates = populate.load_issue_templates(args.data_file)
    session = populate.make_session(args.token)

    print(populate.header(f"Grading data collector - {DELIVERABLE_LABEL}"))
    print(populate.dim(f"  base repo:    {args.owner}/{args.repo}"))
    print(populate.dim(f"  templates:    {os.path.basename(args.data_file)} "
                       f"({len(templates)} assignment issue(s))"))
    print(populate.dim(f"  student_repos: {repos_dir}"))

    print(populate.header(f"\nFetching forks of {args.owner}/{args.repo}..."))
    forks = populate.get_forks(session, args.owner, args.repo)

    users = [u.lower() for u in (args.fork_user or [])]
    if args.test and not users:
        users = [DEFAULT_TEST_FORK_USER.lower()]
    if users:
        forks = [f for f in forks if f["owner"]["login"].lower() in users]
    if args.limit:
        forks = forks[:args.limit]

    if not forks:
        print(populate.warn("No matching forks. Nothing to do."))
        return 1

    print(populate.header(f"Processing {len(forks)} fork(s):"))
    for f in forks:
        print(populate.dim(f"  - {f['full_name']}"))
    if args.dry_run:
        print(populate.info("\n[dry-run] stopping before any clone or API call."))
        return 0

    records: list[dict] = []
    for i, fork in enumerate(forks, 1):
        login = fork["owner"]["login"]
        repo = fork["name"]
        print(populate.header(f"\n[{i}/{len(forks)}] {fork['full_name']}"))

        rec: dict = {
            "fork": {
                "full_name": fork["full_name"],
                "owner": login,
                "name": repo,
                "html_url": fork["html_url"],
                "clone_url": fork["clone_url"],
                "default_branch": fork.get("default_branch", "main"),
                "created_at": fork.get("created_at"),
                "pushed_at": fork.get("pushed_at"),
                "has_issues": fork.get("has_issues", False),
            },
            "collected_at": _now(),
        }

        dest = os.path.join(repos_dir, login)

        # 1 + 2: clone / fetch and read git history
        if not args.skip_git:
            try:
                if not args.skip_clone:
                    dest, action = clone_or_update(fork, repos_dir)
                    rec["git_action"] = action
                    print(populate.success(f"  {action}: {dest}"))
                elif os.path.isdir(os.path.join(dest, ".git")):
                    rec["git_action"] = "reused"
                    print(populate.dim(f"  reused existing clone: {dest}"))
                else:
                    raise RuntimeError("no existing clone and --skip-clone set")
                rec["git"] = collect_git(dest, fork.get("default_branch", "main"))
                rec["git"]["path"] = dest
                g = rec["git"]
                print(populate.info(
                    f"  git: {g['total_commits']} commit(s), {g['distinct_authors']} author(s), "
                    f"{g['active_days']} active day(s), "
                    f"branches +{len(g['extra_branches'])} "
                    f"[{', '.join(f'{k}:{v}' for k, v in sorted(g['kinds'].items()))}]"
                ))
            except (RuntimeError, OSError) as exc:
                rec["git_error"] = str(exc)
                print(populate.error(f"  git failed: {exc}"))

        # README.md presence (checked on the working tree of whatever clone we have)
        if os.path.isdir(os.path.join(dest, ".git")):
            rec["readme_present"] = has_readme(dest)

        # 3: issue data
        if not args.skip_issues:
            if not fork.get("has_issues", False):
                rec["issues_error"] = "issues disabled on fork"
                print(populate.warn("  issues disabled on this fork."))
            else:
                try:
                    rec["issues"], rec["assignment_issues"] = collect_issues(
                        session, login, repo, templates, args.owner
                    )
                    rec["pull_requests"] = collect_prs(session, login, repo)
                    s, a, p = rec["issues"], rec["assignment_issues"], rec["pull_requests"]
                    print(populate.info(
                        f"  issues: {s['count']} total ({s['open']} open / {s['closed']} closed, "
                        f"{s['closed_by_student']} by student), {s['interacted']} interacted, "
                        f"checkboxes {s['checkboxes_done']}/{s['checkboxes_total']} | "
                        f"assignment: {sum(x['found'] for x in a)}/{len(a)} present, "
                        f"{sum(x['interacted'] for x in a)} touched | "
                        f"PRs: {p['count']} ({p['merged']} merged, {p['to_upstream']} to upstream)"
                    ))
                except Exception as exc:  # noqa: BLE001 - keep going to next fork
                    rec["issues_error"] = str(exc)
                    print(populate.error(f"  issues failed: {exc}"))

        # per-fork json
        try:
            os.makedirs(dest, exist_ok=True)
            with open(os.path.join(dest, "grading_data.json"), "w", encoding="utf-8") as fh:
                json.dump(rec, fh, indent=2, ensure_ascii=False)
        except OSError as exc:
            print(populate.warn(f"  could not write grading_data.json: {exc}"))

        records.append(rec)

    # combined outputs
    os.makedirs(out_dir, exist_ok=True)

    # One flat row per fork, in a fixed column order. These are the same
    # rows that go into summary.csv / the grading workbook - kept in
    # summary.json too so every gathered metric is visible there without
    # digging through the nested "forks" records.
    rows = [{col: summary_row(r).get(col, "") for col in SUMMARY_COLUMNS} for r in records]

    payload = {
        "collected_at": _now(),
        "base_repo": f"{args.owner}/{args.repo}",
        "columns": SUMMARY_COLUMNS,
        "summary": rows,
        "forks": records,
    }

    written: list[tuple[str, bool]] = []
    written.append(_save_output(
        os.path.join(out_dir, "summary.json"),
        lambda p: _write_json(p, payload), label="summary.json",
    ))
    written.append(_save_output(
        os.path.join(out_dir, "summary.csv"),
        lambda p: _write_csv(p, rows), label="summary.csv",
    ))
    if not args.no_xlsx:
        try:
            written.append(_save_output(
                os.path.join(out_dir, XLSX_NAME),
                lambda p: write_xlsx(p, records, rows, f"{args.owner}/{args.repo}", _now()),
                label=XLSX_NAME,
            ))
        except ImportError:
            print(populate.warn(
                f"  openpyxl not installed - skipping {XLSX_NAME} "
                "(pip install openpyxl, or pass --no-xlsx)."
            ))

    print(populate.header("\nOutputs:"))
    for p, fallback in written:
        line = f"  {p}"
        print(populate.warn(line + "   (fallback name - original was locked)")
              if fallback else populate.header(line))

    if any(fallback for _, fallback in written):
        print(populate.error(
            "\n  One or more files were open (locked by Excel) and were written under a\n"
            f"  timestamped name instead. Close {XLSX_NAME} / summary.csv and re-run\n"
            "  to refresh the canonical files."
        ))

    print(populate.success(f"Done - {len(records)} fork(s) processed."))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
